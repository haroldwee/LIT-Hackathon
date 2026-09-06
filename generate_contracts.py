import os
import io
import random
import fitz
import docx
from PIL import Image, ImageDraw, ImageFont, ImageFilter

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, 'sample_contracts')
os.makedirs(OUT, exist_ok=True)
random.seed(42)

ARIAL = r'C:\Windows\Fonts\arial.ttf'
ARIAL_BD = r'C:\Windows\Fonts\arialbd.ttf'


def wrap_text(text, width=88):
    lines = []
    for para in text.split('\n'):
        if not para.strip():
            lines.append('')
            continue
        words = para.split()
        cur = ''
        for w in words:
            if len(cur) + len(w) + 1 <= width:
                cur = (cur + ' ' + w).strip()
            else:
                lines.append(cur)
                cur = w
        if cur:
            lines.append(cur)
    return lines


def render_scanned_pdf(title, body, path):
    lines = wrap_text(body, width=86)
    pages = [lines[i:i + 46] for i in range(0, len(lines), 46)]
    pdf = fitz.open()
    font = ImageFont.truetype(ARIAL, 26)
    title_font = ImageFont.truetype(ARIAL_BD, 30)
    for page_lines in pages:
        img = Image.new('RGB', (1700, 2200), (252, 251, 248))
        draw = ImageDraw.Draw(img)
        y = 130
        draw.text((120, 60), title, fill=(10, 10, 10), font=title_font)
        for line in page_lines:
            draw.text((120, y), line, fill=(20, 20, 20), font=font)
            y += 40
        for _ in range(500):
            x = random.randint(0, 1699)
            yy = random.randint(0, 2199)
            g = random.randint(120, 200)
            draw.point((x, yy), fill=(g, g, g))
        for _ in range(3):
            x1 = random.randint(0, 1600)
            y1 = random.randint(0, 2199)
            draw.line([(x1, y1), (x1 + random.randint(20, 80), y1)], fill=(180, 180, 180), width=1)
        img = img.rotate(random.uniform(-1.2, 1.2), expand=False, fillcolor=(252, 251, 248))
        img = img.filter(ImageFilter.GaussianBlur(0.4))
        buf = io.BytesIO()
        img.save(buf, format='PNG')
        page = pdf.new_page(width=595, height=842)
        page.insert_image(page.rect, stream=buf.getvalue())
    pdf.save(path)
    pdf.close()


def make_text_pdf(title, body, path):
    pdf = fitz.open()
    chunks = []
    lines = body.split('\n')
    cur = []
    count = 0
    for line in lines:
        cur.append(line)
        count += len(line) + 1
        if count > 2500:
            chunks.append('\n'.join(cur))
            cur = []
            count = 0
    if cur:
        chunks.append('\n'.join(cur))
    for chunk in chunks:
        page = pdf.new_page(width=595, height=842)
        rect = fitz.Rect(50, 50, 545, 792)
        page.insert_textbox(rect, chunk, fontname='helv', fontsize=9, align=fitz.TEXT_ALIGN_LEFT)
    pdf.save(path)
    pdf.close()


def make_docx(title, body, path):
    d = docx.Document()
    d.add_heading(title, level=1)
    for para in body.split('\n'):
        if para.strip():
            d.add_paragraph(para)
        else:
            d.add_paragraph('')
    d.save(path)


def build_contract(c):
    deliv = '\n'.join(f'    {d}' for d in c['deliverables'])
    body = f"""CONTRACT NO: {c['no']}

This {c['kind']} (the "Agreement") is entered into by and between {c['p1']} ("{c['r1']}") and {c['p2']} ("{c['r2']}").

Effective Date: {c['eff']}
Expiration Date: {c['exp']}

1. SCOPE OF AGREEMENT
{c['scope']}

2. TERM
The Agreement commences on the Effective Date and continues in full force until the Expiration Date, unless terminated earlier in accordance with Section 8. Renewal requires written agreement of both Parties no later than sixty (60) days before the Expiration Date.

3. OBLIGATIONS AND DELIVERABLES
The Parties agree to the following milestones and obligations:
{deliv}
Each Party shall use commercially reasonable efforts to meet every deadline stated above. A deliverable not completed by its due date constitutes a material breach unless an extension is agreed in writing.

4. PAYMENT TERMS
Invoices shall be issued monthly and are payable within thirty (30) days of receipt. Late payments accrue interest at 1.5% per month. {c['payment']}

5. CONFIDENTIALITY
Each Party shall keep all non-public information received under this Agreement strictly confidential and use it solely for the purposes of this Agreement. This obligation survives termination for a period of three (3) years.

6. REPRESENTATIONS AND WARRANTIES
Each Party represents that it has full power and authority to enter into this Agreement and that its performance will not violate any other agreement or applicable law.

7. LIMITATION OF LIABILITY
Except for breaches of confidentiality or gross negligence, neither Party shall be liable for indirect, incidental, or consequential damages. Aggregate liability shall not exceed the total fees paid in the twelve (12) months preceding the claim.

8. TERMINATION
Either Party may terminate this Agreement for material breach upon thirty (30) days written notice if the breach remains uncured. Upon termination, all outstanding deliverables due before the termination date shall be completed and paid.

9. GOVERNING LAW
This Agreement shall be governed by the laws of the State of {c['law']}.

10. MISCELLANEOUS
This Agreement constitutes the entire understanding between the Parties and supersedes all prior discussions. Amendments must be in writing and signed by both Parties. Neither Party may assign this Agreement without prior written consent.

IN WITNESS WHEREOF, the Parties have executed this Agreement as of the Effective Date.

SIGNED for and on behalf of {c['p1']}:

____________________________
{c['s1']}
{c['t1']}
Date: {c['eff']}


SIGNED for and on behalf of {c['p2']}:

____________________________
{c['s2']}
{c['t2']}
Date: {c['eff']}
"""
    return c['title'] + '\n\n' + body


contracts = [
    dict(no='C-2025-001', filename='master_services_agreement_acme.docx', fmt='docx', scanned=False,
         title='MASTER SERVICES AGREEMENT', kind='Master Services Agreement',
         p1='Acme Retail Group Inc.', r1='Customer', p2='NimbusTech Solutions LLC', r2='Provider',
         eff='01/15/2025', exp='01/14/2027', law='Delaware',
         scope='Provider shall deliver managed IT services, including cloud hosting, 24/7 monitoring, and helpdesk support, to Customer as described in Statement(s) of Work executed under this Agreement.',
         payment='Fees are fixed for the first contract year and may be adjusted annually by up to 4%.',
         deliverables=[
             'Deliverable 1 - Cloud environment migration due by 04/15/2025.',
             'Milestone: 10/15/2025 - completion of phase two onboarding.',
             'Deadline: 01/15/2026 - annual service review report.',
         ],
         s1='Jordan Ellis', t1='VP Procurement', s2='Maria Chen', t2='Chief Operating Officer'),

    dict(no='C-2025-002', filename='enterprise_subscription_terms_brightline.pdf', fmt='pdf', scanned=False,
         title='ENTERPRISE SUBSCRIPTION TERMS', kind='Enterprise Subscription Agreement',
         p1='Brightline Software Corp.', r1='Vendor', p2='Harborview Financial Group', r2='Subscriber',
         eff='03/01/2025', exp='02/28/2026', law='New York',
         scope='Vendor grants Subscriber a non-exclusive, non-transferable right to access the Brightline Analytics Platform for up to 500 named users during the Term.',
         payment='Annual subscription fee of $240,000 invoiced upon the Effective Date.',
         deliverables=[
             'Deliverable 1 - user provisioning complete due by 06/01/2025.',
             'Deadline: 12/01/2025 - mid-term business review.',
         ],
         s1='Priya Raman', t1='SVP Sales', s2='Thomas Whitfield', t2='Director of Technology'),

    dict(no='C-2026-003', filename='service_level_agreement_northwind.pdf', fmt='pdf', scanned=False,
         title='SERVICE LEVEL AGREEMENT', kind='Service Level Agreement',
         p1='Northwind Cloud Services Inc.', r1='Supplier', p2='Vertex Health Systems', r2='Customer',
         eff='05/01/2026', exp='10/01/2026', law='California',
         scope='Supplier shall host the Customer patient portal with 99.95% measured uptime, disaster recovery with RPO of 15 minutes and RTO of 4 hours.',
         payment='Monthly service fee of $18,500, subject to service credits of 5% per each 0.1% uptime shortfall.',
         deliverables=[
             'Milestone: 07/01/2026 - failover test report due by this date.',
             'Deliverable 2 - security audit completion due by 09/01/2026.',
         ],
         s1='Alan Brooks', t1='VP Operations', s2='Susan Knight', t2='CIO'),

    dict(no='C-2025-004', filename='raw_components_supply_ironclad.docx', fmt='docx', scanned=False,
         title='RAW COMPONENTS SUPPLY AGREEMENT', kind='Supply Agreement',
         p1='Ironclad Materials Corp.', r1='Supplier', p2='Precision Motors Assemblies LLC', r2='Buyer',
         eff='02/01/2025', exp='01/31/2027', law='Ohio',
         scope='Supplier shall supply aluminum alloy components meeting specification IM-4471 in the quantities set out in rolling quarterly purchase orders.',
         payment='Prices fixed for twelve months; thereafter adjusted per the published LME index.',
         deliverables=[
             'Deliverable 1 - first quarterly shipment due by 05/01/2025.',
             'Milestone: 11/01/2025 - qualification of second production line.',
             'Deadline: 05/01/2026 - annual quality certification.',
         ],
         s1='Viktor Hale', t1='Director of Sales', s2='Erin Walsh', t2='Head of Supply Chain'),

    dict(no='C-2025-005', filename='logistics_services_swiftfreight.docx', fmt='docx', scanned=False,
         title='LOGISTICS SERVICES AGREEMENT', kind='Logistics Services Agreement',
         p1='SwiftFreight Logistics Ltd.', r1='Carrier', p2='Orchard Valley Foods Inc.', r2='Shipper',
         eff='06/01/2025', exp='09/20/2026', law='Illinois',
         scope='Carrier shall provide refrigerated road freight across the Midwest region, including temperature monitoring and proof-of-delivery within 24 hours.',
         payment='Rates per lane as per Schedule A; fuel surcharge adjusted monthly.',
         deliverables=[
             'Deliverable 1 - fleet onboarding due by 07/15/2025.',
             'Milestone: 01/15/2026 - telematics integration live.',
             'Deadline: 09/15/2026 - final cold-chain audit of the term.',
         ],
         s1='Derek Nolan', t1='Regional Manager', s2='Hannah Byrd', t2='Logistics Director'),

    dict(no='C-2024-006', filename='manufacturing_agreement_precisiongears.pdf', fmt='pdf', scanned=False,
         title='MANUFACTURING AGREEMENT', kind='Manufacturing Agreement',
         p1='Precision Gears Manufacturing Inc.', r1='Manufacturer', p2='Apex Robotics Corp.', r2='Purchaser',
         eff='04/01/2024', exp='03/31/2026', law='Massachusetts',
         scope='Manufacturer shall produce and supply planetary gear assemblies per drawings REV-C, with monthly volumes between 800 and 1,200 units.',
         payment='Unit price of $146.50 firm through the Expiration Date; tooling amortized over the first 6 months.',
         deliverables=[
             'Milestone: 08/01/2024 - pilot batch acceptance.',
             'Deliverable 2 - PPAP level 3 documentation due by 11/15/2024.',
             'Deadline: 03/15/2026 - end-of-run final inspection report.',
         ],
         s1='Klaus Meier', t1='Plant Director', s2='Irene Chen', t2='VP Hardware'),

    dict(no='C-2024-007', filename='office_lease_400_main.pdf', fmt='pdf', scanned=False,
         title='COMMERCIAL OFFICE LEASE', kind='Lease Agreement',
         p1='Main Street Property Holdings LLC', r1='Landlord', p2='Crestpoint Consulting Group Inc.', r2='Tenant',
         eff='09/01/2024', exp='08/31/2027', law='New York',
         scope='Landlord leases to Tenant approximately 12,400 square feet on floors 8 and 9 of 400 Main Street for general office use.',
         payment='Base rent of $58,900 per month, plus pro-rata operating expenses, payable on the first business day of each month.',
         deliverables=[
             'Milestone: 09/01/2025 - first annual rent review.',
             'Deliverable 2 - HVAC upgrade completion due by 03/01/2026.',
             'Deadline: 09/01/2026 - second annual rent review.',
         ],
         s1='Gerald Pratt', t1='Managing Member', s2='Fiona Adeyemi', t2='Chief Executive Officer'),

    dict(no='C-2025-008', filename='warehouse_lease_bayside.docx', fmt='docx', scanned=False,
         title='WAREHOUSE LEASE AGREEMENT', kind='Lease Agreement',
         p1='Bayside Logistics Park LP', r1='Lessor', p2='Meridian Distribution Co.', r2='Lessee',
         eff='01/01/2025', exp='12/31/2029', law='New Jersey',
         scope='Lessor leases Warehouse Unit 4 comprising 48,000 square feet with 14 dock doors and 32-foot clear height for distribution operations.',
         payment='Triple net rent of $31,200 per month; Lessee responsible for taxes, insurance, and common area maintenance.',
         deliverables=[
             'Deliverable 1 - racking installation due by 04/01/2025.',
             'Milestone: 01/01/2027 - rent escalation to $32,450 per month.',
             'Deadline: 12/01/2029 - pre-expiration facility condition survey.',
         ],
         s1='Robert Kim', t1='Asset Manager', s2='Angela Cruz', t2='Operations VP'),

    dict(no='C-2025-009', filename='equipment_lease_forklift_fleet.docx', fmt='docx', scanned=False,
         title='EQUIPMENT LEASE AGREEMENT', kind='Lease Agreement',
         p1='LiftSource Equipment Finance Inc.', r1='Lessor', p2='Beacon Building Supply LLC', r2='Lessee',
         eff='07/01/2025', exp='06/30/2027', law='Texas',
         scope='Lessor leases to Lessee twelve Class III electric forklifts and two yard trucks, including scheduled maintenance and parts.',
         payment='Monthly lease payment of $9,850 including maintenance; usage capped at 1,200 hours per unit per year.',
         deliverables=[
             'Deliverable 1 - delivery and operator training due by 08/15/2025.',
             'Milestone: 07/01/2026 - annual inspection certification.',
             'Deadline: 06/15/2027 - return condition inspection.',
         ],
         s1='Paul Gagnon', t1='Account Executive', s2='Dana Whitmore', t2='Branch Manager'),

    dict(no='C-2025-010', filename='mutual_nda_technova.pdf', fmt='pdf', scanned=False,
         title='MUTUAL NON-DISCLOSURE AGREEMENT', kind='Mutual Non-Disclosure Agreement',
         p1='TechNova Labs Inc.', r1='Disclosing Party', p2='Quantum Reach Media LLC', r2='Receiving Party',
         eff='02/15/2025', exp='02/14/2027', law='California',
         scope='The Parties wish to explore a joint product integration and may disclose Confidential Information to each other in connection with that evaluation.',
         payment='No fees payable under this Agreement.',
         deliverables=[
             'Deliverable 1 - exchange of technical documentation due by 05/01/2025.',
             'Deadline: 11/30/2026 - final evaluation summary.',
         ],
         s1='Lena Okafor', t1='CTO', s2='Marcus Reid', t2='Managing Partner'),

    dict(no='C-2025-011', filename='employee_confidentiality_agreement.docx', fmt='docx', scanned=False,
         title='EMPLOYEE CONFIDENTIALITY AND INVENTION ASSIGNMENT AGREEMENT', kind='Employee Confidentiality Agreement',
         p1='Helix Bioworks Inc.', r1='Company', p2='Daniel Mercer', r2='Employee',
         eff='03/10/2025', exp='03/10/2028', law='Washington',
         scope='Employee shall protect Company confidential information, including research data and formulations, and assign inventions made during employment to the Company.',
         payment='Consideration is the salary and benefits payable under the employment agreement.',
         deliverables=[
             'Milestone: 04/10/2025 - completion of compliance training due by this date.',
             'Deadline: 03/10/2026 - annual reaffirmation of obligations.',
         ],
         s1='Nadia Fournier', t1='Chief People Officer', s2='Daniel Mercer', t2='Research Scientist'),

    dict(no='C-2026-012', filename='vendor_nda_cloudnine.pdf', fmt='pdf', scanned=False,
         title='VENDOR NON-DISCLOSURE AGREEMENT', kind='Non-Disclosure Agreement',
         p1='CloudNine Analytics Inc.', r1='Discloser', p2='Stonebridge Insurance Services', r2='Recipient',
         eff='08/01/2026', exp='08/01/2027', law='Illinois',
         scope='Discloser may share product roadmaps and pricing models with Recipient to support a proposed reseller evaluation.',
         payment='No fees payable under this Agreement.',
         deliverables=[
             'Deliverable 1 - initial data room access due by 09/01/2026.',
         ],
         s1='Owen Delgado', t1='VP Partnerships', s2='Grace Lin', t2='Contracts Manager'),

    dict(no='C-2025-013', filename='regional_distribution_pacificrim.pdf', fmt='pdf', scanned=False,
         title='REGIONAL DISTRIBUTION AGREEMENT', kind='Distribution Agreement',
         p1='Pacific Rim Trading Corp.', r1='Distributor', p2='Aurora Home Goods Ltd.', r2='Supplier',
         eff='03/01/2025', exp='02/28/2027', law='Oregon',
         scope='Distributor shall promote, sell, and service Supplier products in Washington, Oregon, and British Columbia, meeting minimum annual purchases of $2.4M.',
         payment='Net 45 terms; 2% discount for payment within 15 days.',
         deliverables=[
             'Milestone: 09/30/2025 - launch in 40 retail doors.',
             'Deliverable 2 - first annual market plan due by 03/31/2026.',
             'Deadline: 09/30/2026 - second annual market plan.',
         ],
         s1='Victor Ito', t1='President', s2='Charlotte May', t2='Global Sales Director'),

    dict(no='C-2026-014', filename='exclusive_distributor_euromarket.docx', fmt='docx', scanned=False,
         title='EXCLUSIVE DISTRIBUTOR AGREEMENT', kind='Exclusive Distribution Agreement',
         p1='EuroMarket Handels GmbH', r1='Exclusive Distributor', p2='Silverleaf Beverages Inc.', r2='Supplier',
         eff='01/01/2026', exp='12/31/2028', law='Hamburg, Germany',
         scope='Supplier grants Distributor exclusive rights to import and distribute Silverleaf products in Germany and Austria, subject to minimum purchase commitments.',
         payment='Prices per Annex B; annual rebate of 3% if purchases exceed EUR 4.5M.',
         deliverables=[
             'Deliverable 1 - regulatory registrations due by 06/30/2026.',
             'Milestone: 01/15/2027 - first annual volume commitment review.',
             'Deadline: 06/30/2027 - mid-term compliance audit.',
         ],
         s1='Hans Keller', t1='Geschaeftsfuehrer', s2='Rebecca Stone', t2='VP International'),

    dict(no='C-2025-015', filename='channel_partner_southerncross.pdf', fmt='pdf', scanned=False,
         title='CHANNEL PARTNER AGREEMENT', kind='Channel Partner Agreement',
         p1='Southern Cross Retail Partners Pty Ltd', r1='Partner', p2='VoltEdge Electronics Inc.', r2='Principal',
         eff='10/01/2025', exp='09/30/2026', law='Victoria, Australia',
         scope='Partner shall resell Principal consumer electronics lines through its retail network and e-commerce channels, adhering to brand guidelines.',
         payment='Wholesale pricing per Schedule 2; quarterly volume rebates up to 5%.',
         deliverables=[
             'Deliverable 1 - store rollout across 25 locations due by 01/01/2026.',
             'Milestone: 07/01/2026 - mid-term sales performance review.',
             'Deadline: 09/15/2026 - final sell-through report for the term.',
         ],
         s1='Bruce Halloran', t1='Managing Director', s2='Nina Patel', t2='Channel Director'),

    dict(no='C-2026-016', filename='nda_meridian_scanned.pdf', fmt='pdf', scanned=True,
         title='NON-DISCLOSURE AGREEMENT (SCANNED)', kind='Non-Disclosure Agreement',
         p1='Meridian Robotics Inc.', r1='Disclosing Party', p2='Falcon Manufacturing LLC', r2='Receiving Party',
         eff='05/05/2026', exp='05/04/2028', law='Michigan',
         scope='Receiving Party may receive proprietary designs and tooling specifications to prepare a quotation for contract manufacturing services.',
         payment='No fees payable under this Agreement.',
         deliverables=[
             'Deliverable 1 - quotation package due by 07/05/2026.',
         ],
         s1='Sofia Grant', t1='VP Engineering', s2='Carl Jensen', t2='Estimating Manager'),

    dict(no='C-2025-017', filename='supplier_contract_vanguard_scanned.pdf', fmt='pdf', scanned=True,
         title='SUPPLIER AGREEMENT (SCANNED)', kind='Supplier Contract',
         p1='Vanguard Packaging Corp.', r1='Supplier', p2='Willowbrook Pharmaceuticals Inc.', r2='Buyer',
         eff='11/01/2025', exp='10/31/2026', law='Pennsylvania',
         scope='Supplier shall furnish corrugated packaging and printed inserts to GMP standards, with lot traceability and certificates of analysis per shipment.',
         payment='Prices per Schedule A valid through the Expiration Date; net 30 payment terms.',
         deliverables=[
             'Deliverable 1 - artwork approval cycle complete due by 12/15/2025.',
             'Milestone: 04/30/2026 - GMP re-audit.',
             'Deadline: 10/15/2026 - end of term reconciliation report.',
         ],
         s1='Miguel Torres', t1='National Account Manager', s2='Janet Fowler', t2='Procurement Lead'),

    dict(no='C-2025-018', filename='lease_harborside_scanned.pdf', fmt='pdf', scanned=True,
         title='LEASE AGREEMENT (SCANNED)', kind='Lease Agreement',
         p1='Harborside Realty Trust', r1='Landlord', p2='Tidewater Marine Services LLC', r2='Tenant',
         eff='12/01/2025', exp='11/30/2027', law='Florida',
         scope='Landlord leases Pier Shed B comprising 9,600 square feet including office mezzanine for marine service operations.',
         payment='Monthly rent of $21,750 plus CAM charges estimated at $2,300 per month.',
         deliverables=[
             'Milestone: 03/01/2026 - dock crane installation due by this date.',
             'Deadline: 11/15/2027 - pre-termination walkthrough.',
         ],
         s1='Elena Vasquez', t1='Portfolio Manager', s2='Peter Lang', t2='General Manager'),

    dict(no='C-2026-019', filename='customer_terms_zenith_scanned.pdf', fmt='pdf', scanned=True,
         title='CUSTOMER TERMS OF SERVICE (SCANNED)', kind='Customer Terms Agreement',
         p1='Zenith Facilities Management Inc.', r1='Service Provider', p2='Grandview Hospitality Group', r2='Customer',
         eff='06/15/2026', exp='06/14/2027', law='Nevada',
         scope='Provider shall deliver janitorial and maintenance services across four hotel properties, five nights per week, per specification sheets S-1 to S-6.',
         payment='Monthly fee of $46,000; additional services billed at schedule rates.',
         deliverables=[
             'Deliverable 1 - site mobilization due by 07/15/2026.',
             'Deadline: 09/30/2026 - first quarterly quality audit.',
         ],
         s1='Alicia Fontaine', t1='Regional Director', s2='Michael Osei', t2='VP Hotel Operations'),

    dict(no='C-2026-020', filename='distribution_agreement_atlas_scanned.pdf', fmt='pdf', scanned=True,
         title='DISTRIBUTION AGREEMENT (SCANNED)', kind='Distribution Agreement',
         p1='Atlas Outdoor Gear Inc.', r1='Supplier', p2='Summit Trail Distributors Ltd.', r2='Distributor',
         eff='08/15/2026', exp='08/14/2028', law='Colorado',
         scope='Distributor shall distribute Supplier products across garden route territories, with minimum annual purchases of $1.8M and quarterly stock reporting.',
         payment='Net 60 terms from invoice date; freight FOB Supplier warehouse.',
         deliverables=[
             'Deliverable 1 - initial inventory order due by 09/30/2026.',
             'Milestone: 02/15/2027 - spring product line launch.',
             'Deadline: 08/01/2027 - annual volume reconciliation.',
         ],
         s1='Rachel Donnelly', t1='Sales Director', s2='Simon Baxter', t2='Managing Director'),
]

for c in contracts:
    body = build_contract(c)
    path = os.path.join(OUT, c['filename'])
    if c['scanned']:
        render_scanned_pdf(c['title'], body, path)
    elif c['fmt'] == 'docx':
        make_docx(c['title'], body, path)
    else:
        make_text_pdf(c['title'], body, path)
    print(f"Created {c['filename']} (scanned={c['scanned']})")

print(f"\nTotal: {len(contracts)} contracts in {OUT}")
