import os
import re
import io
import json
import uuid
import base64
import datetime
import traceback
from pathlib import Path
from datetime import datetime as dt
from urllib.parse import urlparse

import requests
from flask import Flask, request, jsonify, send_from_directory, send_file
from flask_cors import CORS
import pymupdf as fitz
import docx
try:
    import pytesseract
    _TESSERACT_CANDIDATES = [
        r'C:\Program Files\Tesseract-OCR\tesseract.exe',
        r'C:\Program Files (x86)\Tesseract-OCR\tesseract.exe',
        os.path.join(os.environ.get('LOCALAPPDATA', ''), r'Programs\Tesseract-OCR\tesseract.exe'),
    ]
    for _tess in _TESSERACT_CANDIDATES:
        if _tess and os.path.isfile(_tess):
            pytesseract.pytesseract.tesseract_cmd = _tess
            break
    pytesseract.get_tesseract_version()
    HAS_TESSERACT = True
except Exception:
    HAS_TESSERACT = False
from PIL import Image, ImageDraw
import numpy as np

app = Flask(__name__)
CORS(app)

UPLOAD_FOLDER = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'uploads')
PROCESSED_FOLDER = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'processed')
DB_FILE = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'contracts.json')

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(PROCESSED_FOLDER, exist_ok=True)

ALLOWED_EXTENSIONS = {'.docx', '.pdf', '.txt'}

CONFIG_FILE = os.path.join(os.path.dirname(__file__), 'config.json')

def load_config():
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, 'r') as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

CONFIG = load_config()
OPENROUTER_API_KEY = CONFIG.get('openrouter_api_key') or os.environ.get('OPENROUTER_API_KEY', '')
OPENROUTER_MODEL = CONFIG.get('openrouter_model', 'google/gemini-2.5-flash')
OPENROUTER_URL = 'https://openrouter.ai/api/v1/chat/completions'
VALID_CONTRACT_TYPES = {'NDAs', 'Supplier Contract', 'Customer Terms', 'Lease', 'Distribution Agreement', 'Agreement'}

def load_contracts():
    if os.path.exists(DB_FILE):
        with open(DB_FILE, 'r') as f:
            return json.load(f)
    return []

def save_contracts(contracts):
    with open(DB_FILE, 'w') as f:
        json.dump(contracts, f, indent=2, default=str)

def allowed_file(filename):
    return os.path.splitext(filename)[1].lower() in ALLOWED_EXTENSIONS

def _number_paragraphs(text):
    blocks = [b.strip() for b in re.split(r'\n\s*\n', text) if b.strip()]
    return '\n\n'.join(f'[[P{j}]]\n{b}' for j, b in enumerate(blocks, start=1))

def extract_text_pages(pdf_path):
    doc = fitz.open(pdf_path)
    pages = []
    try:
        for page in doc:
            text = page.get_text()
            if not text.strip() and HAS_TESSERACT:
                pix = page.get_pixmap(dpi=300)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                text = pytesseract.image_to_string(img)
            pages.append(_number_paragraphs(text))
    finally:
        doc.close()
    return pages

def extract_text_from_pdf(pdf_path):
    pages = extract_text_pages(pdf_path)
    return '\n\n'.join(f'[[PAGE {i}]]\n{t}' for i, t in enumerate(pages, start=1))

def extract_text_from_docx(docx_path):
    doc = docx.Document(docx_path)
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    full_text = _number_paragraphs('\n\n'.join(paras))
    for table in doc.tables:
        for row in table.rows:
            row_text = "\t".join([cell.text for cell in row.cells])
            full_text += "\n" + row_text
    return full_text

def extract_dates(text):
    date_patterns = [
        r'(?:effective|commence|start|begin|dated)\s*(?:as\s+of|on)?\s*:?\s*(\d{1,2}[\/\-\.]\d{1,2}[\/\-\.]\d{2,4})',
        r'(\d{1,2}[\/\-\.]\d{1,2}[\/\-\.]\d{2,4})\s*(?:to|through|until|ending|ending\s*on)',
        r'(\d{1,2}[\/\-\.]\d{1,2}[\/\-\.]\d{2,4})\s*-?\s*(\d{1,2}[\/\-\.]\d{1,2}[\/\-\.]\d{2,4})',
        r'(?:effective\s*date|commencement\s*date|start\s*date|beginning\s*date)\s*:?\s*(\d{1,2}[\/\-\.]\d{1,2}[\/\-\.]\d{2,4})',
        r'(?:expiration|expiry|termination|end)\s*(?:date|on)\s*:?\s*(\d{1,2}[\/\-\.]\d{1,2}[\/\-\.]\d{2,4})',
        r'(\d{1,2}\s+(?:January|February|March|April|May|June|July|August|September|October|November|December|Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\w*\s+\d{4})',
        r'(\d{4}-\d{2}-\d{2})',
        r'(\d{2}/\d{2}/\d{4})',
        r'(\d{1,2}\s+\w+\s+\d{4})',
    ]
    dates_found = []
    for pattern in date_patterns:
        matches = re.findall(pattern, text, re.IGNORECASE)
        for m in matches:
            if isinstance(m, tuple):
                dates_found.extend([d for d in m if d])
            else:
                dates_found.append(m)
    return dates_found

def parse_date(date_str):
    for fmt in ['%m/%d/%Y', '%m/%d/%y', '%d/%m/%Y', '%d/%m/%y', '%Y-%m-%d', '%B %d, %Y', '%b %d, %Y', '%d %B %Y', '%d %b %Y']:
        try:
            return dt.strptime(date_str.strip(), fmt)
        except ValueError:
            continue
    return None

EXTRACTION_PROMPT = """You are a legal contract analyst. Extract structured data from the contract document provided.

FIELD RULES:
- contract_name: the official title of the agreement. Use null only if truly unidentifiable.
- contract_type: classify as exactly one of "NDAs", "Supplier Contract", "Customer Terms", "Lease", "Distribution Agreement", or "Agreement" if none fit.
- parties: every entity entering the agreement. "role" is their label in the document (e.g. "Customer", "Supplier", "Landlord") or null.
- start_date / end_date: YYYY-MM-DD format, or null. null for perpetual agreements.
- deliverables: every dated obligation - deliverables, milestones, deadlines, reviews, audits, rent reviews, inspections, reports, orders. "description" is a short label for what is due. Do not invent dates. Empty list if none.
- facts: the key commercial/legal terms. ALWAYS include these four core facts first: "Contract Type", "Parties", "Start Date", "End Date". Then add every significant term present, such as: Payment Terms, Termination, Renewal, Confidentiality, Liability, Indemnity, Governing Law, Exclusivity, Minimum Commitment, Notice Period, Intellectual Property, Dispute Resolution. Omit terms not in the document. "value" is a concise summary (1-2 sentences). "confidence": "high" if explicit and unambiguous, "medium" if stated but needs interpretation, "low" if unclear.

CITATION RULES (apply to EVERY field, every party, every deliverable, every fact):
- The document text contains page markers like [[PAGE 4]] and paragraph markers like [[P3]] (meaning paragraph 3 of the current page). Set "page" to the number of the nearest [[PAGE n]] marker at or before the supporting text, and "paragraph" to the number of the nearest [[Pn]] marker at or before the supporting text. If the document was provided as page images, each image is labelled "Page N:" - use that number for "page" and null for "paragraph". Otherwise use null.
- "clause": the closest clause/section identifier or heading at or above the supporting text (e.g. "Section 5.2", "Clause 4", "Preamble", "Title", "Signature Block"). null if none.
- "quote": a short verbatim quote from the document (max ~25 words) that supports the value, exactly as it appears.
- "basis": "found" when the value is explicitly stated in the document and the quote directly supports it; "inferred" when you deduced, calculated, classified or assumed the value. Never claim "found" without a real supporting quote.

Respond with ONLY a JSON object, no markdown fences, no explanation, in exactly this schema:
{"contract_name": {"value": string | null, "basis": string, "page": int | null, "paragraph": int | null, "clause": string | null, "quote": string},
 "contract_type": {"value": string, "basis": string, "page": int | null, "paragraph": int | null, "clause": string | null, "quote": string},
 "parties": [{"name": string, "role": string | null, "basis": string, "page": int | null, "paragraph": int | null, "clause": string | null, "quote": string}],
 "start_date": {"value": string | null, "basis": string, "page": int | null, "paragraph": int | null, "clause": string | null, "quote": string},
 "end_date": {"value": string | null, "basis": string, "page": int | null, "paragraph": int | null, "clause": string | null, "quote": string},
 "deliverables": [{"date": string, "description": string, "basis": string, "page": int | null, "paragraph": int | null, "clause": string | null, "quote": string}],
 "facts": [{"term": string, "value": string, "confidence": string, "basis": string, "page": int | null, "paragraph": int | null, "clause": string | null, "quote": string}]}"""

def _call_openrouter(messages, max_tokens=4000):
    headers = {
        'Authorization': f'Bearer {OPENROUTER_API_KEY}',
        'Content-Type': 'application/json',
        'HTTP-Referer': 'http://localhost:5000',
        'X-Title': 'Contract Manager'
    }
    payload = {
        'model': OPENROUTER_MODEL,
        'messages': messages,
        'temperature': 0,
        'max_tokens': max_tokens
    }
    r = requests.post(OPENROUTER_URL, headers=headers, json=payload, timeout=120)
    r.raise_for_status()
    data = r.json()
    return data['choices'][0]['message']['content']

def _parse_json_content(content):
    content = content.strip()
    if content.startswith('```'):
        content = re.sub(r'^```(?:json)?\s*', '', content)
        content = re.sub(r'\s*```$', '', content).strip()
    start, end = content.find('{'), content.rfind('}')
    if start == -1 or end == -1:
        raise ValueError('No JSON object found in AI response')
    return json.loads(content[start:end + 1])

def _normalize_date(v):
    if not v:
        return None
    v = str(v).strip()
    if not v or v.lower() in ('null', 'none', 'n/a', 'perpetual', 'unknown'):
        return None
    try:
        return dt.strptime(v, '%Y-%m-%d').date().isoformat()
    except ValueError:
        parsed = parse_date(v)
        return parsed.date().isoformat() if parsed else None

def _norm_citation(page, paragraph, clause, quote):
    try:
        p = int(page)
        if p < 1:
            p = None
    except (TypeError, ValueError):
        p = None
    try:
        para = int(paragraph)
        if para < 1:
            para = None
    except (TypeError, ValueError):
        para = None
    q = ' '.join(str(quote or '').split())[:300]
    return {
        'page': p,
        'paragraph': para,
        'clause': (str(clause).strip()[:80] or None) if clause else None,
        'quote': q
    }

def _norm_basis(basis, citation):
    b = str(basis or '').lower().strip()
    if b in ('found', 'inferred'):
        return b
    return 'found' if (citation and citation.get('quote')) else 'inferred'

def _norm_scalar_field(v, fallback=None):
    if isinstance(v, dict):
        val = v.get('value')
        cit = _norm_citation(v.get('page'), v.get('paragraph'), v.get('clause'), v.get('quote'))
        basis = _norm_basis(v.get('basis'), cit)
    else:
        val = v
        cit = {'page': None, 'paragraph': None, 'clause': None, 'quote': ''}
        basis = 'inferred'
    if isinstance(val, str):
        val = val.strip()
        if val.lower() in ('null', 'none', 'n/a', ''):
            val = None
    if val is None and fallback:
        val = fallback
    return val, cit, basis

def _normalize_ai_result(result, filename):
    fallback_name = os.path.splitext(filename)[0].replace('_', ' ').replace('-', ' ').title()

    name, name_cit, name_basis = _norm_scalar_field(result.get('contract_name'), fallback=fallback_name)
    ctype_raw, ctype_cit, ctype_basis = _norm_scalar_field(result.get('contract_type'))
    ctype = ctype_raw if ctype_raw in VALID_CONTRACT_TYPES else None

    parties = []
    party_details = []
    parties_cit = None
    for p in (result.get('parties') or []):
        if isinstance(p, dict):
            pname = ' '.join(str(p.get('name') or '').split())
            pcit = _norm_citation(p.get('page'), p.get('paragraph'), p.get('clause'), p.get('quote'))
            pbasis = _norm_basis(p.get('basis'), pcit)
            prole = p.get('role')
            prole = ' '.join(str(prole).split())[:60] if prole else None
        else:
            pname = ' '.join(str(p or '').split())
            pcit = {'page': None, 'paragraph': None, 'clause': None, 'quote': ''}
            pbasis = 'inferred'
            prole = None
        if not pname or pname.lower() in [x.lower() for x in parties]:
            continue
        parties.append(pname)
        party_details.append({'name': pname, 'role': prole, 'basis': pbasis, **pcit})
        if parties_cit is None and pcit.get('quote'):
            parties_cit = pcit
    parties = parties[:4]
    party_details = party_details[:4]

    start_raw, start_cit, start_basis = _norm_scalar_field(result.get('start_date'))
    start_iso = _normalize_date(start_raw)
    end_raw, end_cit, end_basis = _norm_scalar_field(result.get('end_date'))
    end_iso = _normalize_date(end_raw)

    deliverables = []
    deliverable_details = []
    deliv_cit = None
    for d in (result.get('deliverables') or []):
        if isinstance(d, dict):
            nd = _normalize_date(d.get('date'))
            dcit = _norm_citation(d.get('page'), d.get('paragraph'), d.get('clause'), d.get('quote'))
            dbasis = _norm_basis(d.get('basis'), dcit)
            desc = ' '.join(str(d.get('description') or 'Deliverable deadline').split())[:150]
        else:
            nd = _normalize_date(d)
            dcit = {'page': None, 'paragraph': None, 'clause': None, 'quote': ''}
            dbasis = 'inferred'
            desc = 'Deliverable deadline'
        if not nd or nd in deliverables:
            continue
        deliverables.append(nd)
        deliverable_details.append({'date': nd, 'description': desc, 'basis': dbasis, **dcit})
        if deliv_cit is None:
            deliv_cit = dcit
    deliverables.sort()
    deliverable_details.sort(key=lambda x: x['date'])

    facts = []
    for f in (result.get('facts') or []):
        if not isinstance(f, dict) or not f.get('term') or not f.get('value'):
            continue
        conf = str(f.get('confidence', 'low')).lower().strip()
        if conf not in ('high', 'medium', 'low'):
            conf = 'low'
        fcit = _norm_citation(f.get('page'), f.get('paragraph'), f.get('clause'), f.get('quote') or f.get('source'))
        facts.append({
            'term': str(f['term']).strip()[:100],
            'value': ' '.join(str(f['value']).split())[:500],
            'confidence': conf,
            'basis': _norm_basis(f.get('basis'), fcit),
            'page': fcit['page'],
            'paragraph': fcit['paragraph'],
            'clause': fcit['clause'],
            'quote': fcit['quote']
        })
    facts = facts[:15]

    core = {'contract type', 'parties', 'start date', 'end date'}
    have = {f['term'].lower() for f in facts}
    core_facts = [
        ('Contract Type', ctype, ctype_cit, ctype_basis, 'high'),
        ('Parties', ', '.join(parties), parties_cit, 'found' if parties_cit else 'inferred', 'high'),
        ('Start Date', start_iso, start_cit, start_basis, 'high'),
        ('End Date', end_iso, end_cit, end_basis, 'high'),
    ]
    for label, val, cit, basis, conf in core_facts:
        if not val or label.lower() in have:
            continue
        facts.append({
            'term': label, 'value': str(val), 'confidence': conf, 'basis': basis,
            'page': cit['page'] if cit else None, 'paragraph': cit['paragraph'] if cit else None, 'clause': cit['clause'] if cit else None,
            'quote': cit['quote'] if cit else ''
        })

    field_citations = {
        'contract_name': {**name_cit, 'basis': name_basis},
        'contract_type': {**ctype_cit, 'basis': ctype_basis},
        'parties': {**(parties_cit or {'page': None, 'paragraph': None, 'clause': None, 'quote': ''}), 'basis': 'found' if parties_cit else 'inferred'},
        'start_date': {**start_cit, 'basis': start_basis},
        'end_date': {**end_cit, 'basis': end_basis},
        'deliverables': {**(deliv_cit or {'page': None, 'paragraph': None, 'clause': None, 'quote': ''}), 'basis': 'found' if deliv_cit else 'inferred'}
    }

    return {
        'contract_name': str(name).strip(),
        'contract_type': ctype,
        'parties': parties,
        'party_details': party_details,
        'start_date': start_iso,
        'end_date': end_iso,
        'deliverables': deliverables,
        'deliverable_details': deliverable_details,
        'field_citations': field_citations,
        'facts': facts
    }

def ai_extract_from_text(text, filename):
    if not OPENROUTER_API_KEY or not text.strip():
        return None
    t = text.strip()
    if len(t) > 30000:
        t = t[:20000] + '\n\n[...middle of document truncated...]\n\n' + t[-8000:]
    messages = [
        {'role': 'system', 'content': EXTRACTION_PROMPT},
        {'role': 'user', 'content': f'Filename: {filename}\n\nDocument text:\n\n{t}'}
    ]
    try:
        content = _call_openrouter(messages)
        return _parse_json_content(content)
    except Exception:
        print(f'AI text extraction failed for {filename}:')
        traceback.print_exc()
        return None

def pdf_image_pages(pdf_path, max_pages=4):
    images = []
    doc = fitz.open(pdf_path)
    try:
        for i, page in enumerate(doc, start=1):
            if len(images) >= max_pages:
                break
            if page.get_text().strip():
                continue
            pix = page.get_pixmap(dpi=150)
            img = Image.frombytes('RGB', [pix.width, pix.height], pix.samples)
            buf = io.BytesIO()
            img.save(buf, format='JPEG', quality=80)
            images.append((i, base64.b64encode(buf.getvalue()).decode('ascii')))
    finally:
        doc.close()
    return images

def ai_extract_from_images(images, filename):
    if not OPENROUTER_API_KEY or not images:
        return None
    page_list = ', '.join(str(n) for n, _ in images)
    content = [{'type': 'text', 'text': f'Filename: {filename}\n\nThese are scanned pages of a contract document, in order: pages {page_list}. Each image is labelled "Page N:" - use that number as the "page" citation. Extract the structured data as instructed.'}]
    for n, b64 in images:
        content.append({'type': 'text', 'text': f'Page {n}:'})
        content.append({'type': 'image_url', 'image_url': {'url': f'data:image/jpeg;base64,{b64}'}})
    messages = [
        {'role': 'system', 'content': EXTRACTION_PROMPT},
        {'role': 'user', 'content': content}
    ]
    try:
        return _parse_json_content(_call_openrouter(messages, max_tokens=4000))
    except Exception:
        print(f'AI vision extraction failed for {filename}:')
        traceback.print_exc()
        return None

def _regex_facts(info):
    facts = []
    if info.get('contract_type'):
        facts.append({'term': 'Contract Type', 'value': info['contract_type'], 'confidence': 'medium', 'basis': 'inferred', 'page': None, 'paragraph': None, 'clause': None, 'quote': ''})
    if info.get('parties'):
        facts.append({'term': 'Parties', 'value': ', '.join(info['parties']), 'confidence': 'medium', 'basis': 'found', 'page': None, 'paragraph': None, 'clause': None, 'quote': ''})
    if info.get('start_date'):
        facts.append({'term': 'Start Date', 'value': info['start_date'][:10], 'confidence': 'medium', 'basis': 'found', 'page': None, 'paragraph': None, 'clause': None, 'quote': ''})
    if info.get('end_date'):
        facts.append({'term': 'End Date', 'value': info['end_date'][:10], 'confidence': 'medium', 'basis': 'found', 'page': None, 'paragraph': None, 'clause': None, 'quote': ''})
    if info.get('deliverables'):
        facts.append({'term': 'Deliverable Deadlines', 'value': ', '.join(d[:10] for d in info['deliverables']), 'confidence': 'medium', 'basis': 'found', 'page': None, 'paragraph': None, 'clause': None, 'quote': ''})
    return facts

def _regex_citations(info):
    nocite = {'page': None, 'paragraph': None, 'clause': None, 'quote': '', 'basis': 'inferred'}
    return {
        'contract_name': dict(nocite),
        'contract_type': dict(nocite),
        'parties': {**nocite, 'basis': 'found' if info.get('parties') else 'inferred'},
        'start_date': {**nocite, 'basis': 'found' if info.get('start_date') else 'inferred'},
        'end_date': {**nocite, 'basis': 'found' if info.get('end_date') else 'inferred'},
        'deliverables': {**nocite, 'basis': 'found' if info.get('deliverables') else 'inferred'}
    }

def process_document(filepath, original_filename):
    ext = os.path.splitext(original_filename)[1].lower()
    text = ''
    if ext == '.pdf':
        text = extract_text_from_pdf(filepath)
    elif ext == '.docx':
        text = extract_text_from_docx(filepath)
    elif ext == '.txt':
        with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
            text = f.read()

    info = None
    method = 'regex'
    if len(text.strip()) < 100 and ext == '.pdf':
        result = ai_extract_from_images(pdf_image_pages(filepath), original_filename)
        if result:
            info = _normalize_ai_result(result, original_filename)
            method = 'ai_vision'
    else:
        result = ai_extract_from_text(text, original_filename)
        if result:
            info = _normalize_ai_result(result, original_filename)
            method = 'ai'

    if info is None:
        if not text.strip():
            text = 'Contract document uploaded. Content extraction pending.'
        info = extract_contract_info(text, original_filename)
        info['facts'] = _regex_facts(info)
        info['party_details'] = [{'name': p, 'role': None, 'basis': 'found', 'page': None, 'paragraph': None, 'clause': None, 'quote': ''} for p in info.get('parties', [])]
        info['deliverable_details'] = [{'date': d[:10], 'description': 'Deliverable deadline', 'basis': 'found', 'page': None, 'paragraph': None, 'clause': None, 'quote': ''} for d in info.get('deliverables', [])]
        info['field_citations'] = _regex_citations(info)

    if info.get('contract_type') is None:
        head_extract = extract_contract_info(text or '', original_filename)
        info['contract_type'] = head_extract['contract_type']

    info['id'] = str(uuid.uuid4())
    info['filename'] = original_filename
    info['uploaded_at'] = dt.now().isoformat()
    info['extraction_method'] = method
    info['text'] = text[:2000] if text else ''
    return info

CANON_TOPICS = [
    'Parties', 'Term', 'Renewal Mechanics & Notice Periods', 'Termination Rights',
    'Payment Obligations', 'Liability Caps', 'Exclusivity / Restrictive Covenants'
]

ALLOWED_REGULATOR_HOSTS = ('sso.agc.gov.sg', 'mom.gov.sg', 'iras.gov.sg', 'acra.gov.sg', 'pdpc.gov.sg')

MATRIX_PROMPT = """You are a Singapore-qualified contract analyst. Build a two-party terms matrix from the contract document provided, with rigorous conflict checking.

IDENTIFY PARTIES: "party1" is the FIRST-named party in the preamble, "party2" is the SECOND-named party.

Build EXACTLY these rows, in this order:
1. "Parties" - who the parties are and the role each plays under the agreement
2. "Term" - commencement date, expiration/end date or perpetual status, term length
3. "Renewal Mechanics & Notice Periods" - renewal process, renewal deadlines, notice periods required
4. "Termination Rights" - who may terminate, on what grounds, notice required, cure periods
5. "Payment Obligations" - who pays whom, amounts/rates, invoicing cadence, payment due periods, late interest
6. "Liability Caps" - exclusions, monetary caps, carve-outs, and who benefits from them
7. "Exclusivity / Restrictive Covenants" - exclusivity grants, non-compete, non-solicitation, assignment restrictions. When exclusivity exists, analyse its precise scope: duration/dates, territory, product/service scope, sales volumes or minimum purchases, channels covered, and any carve-outs. Write "None stated" if the contract is silent.

ROW VALUES:
- "party1" / "party2": the position, obligation or right SPECIFIC to that party for this topic (concise, 1-2 sentences). null if the topic imposes nothing specific on that party.
- "both": if the information applies EQUALLY to both parties, put the shared value here (and set party1/party2 to null). Use this for mutual obligations like confidentiality, governing law, dispute resolution.
- "confidence": "high" when the contract language is unambiguous and clearly answers the topic; "low" when (a) the wording is ambiguous, vague, one-sided, incomplete or self-contradictory, OR (b) the term may conflict with Singapore legislation or published guidance from MOM, IRAS, ACRA or PDPC, OR (c) it conflicts with one of the OTHER CONTRACTS listed below, OR (d) it is internally inconsistent with another clause of THIS contract (see INTERNAL CONSISTENCY CHECK).
- "low_reason": REQUIRED whenever confidence is "low" (null otherwise):
    - "type": "ambiguity" | "statute" | "conflict" | "internal"
    - "explanation": for ambiguity - precisely why the wording is unclear or open to multiple interpretations, naming each possible reading; for statute - the EXACT discrepancy between the clause and the statutory requirement, QUOTING the statutory provision text alongside the contract text; for conflict - what the other contract says that clashes, quoting its reported terms; for internal - what the two clauses say that contradicts each other
    - "recommendation": concrete actionable steps - what to renegotiate or clarify with the counterparty, specific wording to add or amend
    - "statute": for statute type only - {"name": act name, "section": "Section X", "url": link. Prefer an exact Singapore Statutes Online act page (e.g. https://sso.agc.gov.sg/Act/PDPA2012, https://sso.agc.gov.sg/Act/CompAct2004, https://sso.agc.gov.sg/Act/UCATA1977) if you are confident it exists; otherwise the relevant portal home: SSO https://sso.agc.gov.sg/ or Search https://sso.agc.gov.sg/Search?Phrase=<keywords>, MOM tripartite guidelines https://www.mom.gov.sg/employment-practices/tripartism-in-singapore/tripartite-guidelines-and-advisories, IRAS e-tax guides https://www.iras.gov.sg/quick-links/e-tax-guides, ACRA regulations https://www.acra.gov.sg/regulations/, PDPC guidelines https://www.pdpc.gov.sg/guidelines-and-consultation}
    - "conflict_with": for conflict type only - {"contract": EXACT contract name copied from the OTHER CONTRACTS list, "detail": what that contract says that clashes, including its quoted terms}
    - "internal_conflict": for internal type only - {"clause": section identifier of the contradicting clause in THIS contract (e.g. "Section 2. TERM"), "page": its page number from [[PAGE n]] markers, "paragraph": its paragraph number from [[Pn]] markers, "quote": short verbatim quote of the contradicting clause}

CROSS-CONTRACT EXCLUSIVITY CHECK: for the Exclusivity / Restrictive Covenants row (and any row where exclusivity matters), compare this contract's exclusivity scope (dates, territory, sales, products, channels) against the exclusivity and restrictive terms reported for the OTHER CONTRACTS. If two contracts grant exclusive or restrictive rights that cannot coexist (overlapping territory, overlapping period, conflicting channel or product exclusivity, minimum purchase promises to different counterparties that cannot both be met), flag the affected rows "low" with type "conflict".

INTERNAL CONSISTENCY CHECK: identify terms DEFINED in the contract (definitions section or italicised/quoted defined terms) and compare each definition against how the term is actually used in the operating clauses. Also check for contradictions between clauses (e.g. the Term clause vs dates stated in operating clauses; a delivery deadline in an operational clause vs the notice periods in the Term clause; differing definitions of the same defined term). When an inconsistency affects a row, set that row "low" with type "internal" and populate "internal_conflict" pointing at the contradicting clause.

SINGAPORE LAW CHECK (flag as statute low-confidence ONLY where genuinely relevant):
- Unfair Contract Terms Act 1977 - unreasonable exclusion/restriction of liability
- Personal Data Protection Act 2012 - confidentiality clauses handling personal data (PDPC advisories)
- Competition Act 2004 s 34 / s 47 - non-competes and exclusivity with anti-competitive effect
- Employment Act 1968 - clauses engaging individuals in employment-like arrangements (MOM tripartite guidelines)
- Income Tax Act 1947 - tax indemnities or gross-up clauses (IRAS)
- Companies Act 1967 - corporate authority matters (ACRA)
Only cite provisions you are confident exist. If unsure of a section number, name the act without a section.

OTHER CONTRACTS IN THE DATABASE (for conflict checking - the list may be empty):
{OTHERS}

Respond with ONLY a JSON object, no markdown fences, in exactly this schema:
{"party1": {"name": string, "role": string | null}, "party2": {"name": string, "role": string | null},
 "rows": [{"topic": string, "party1": string | null, "party2": string | null, "both": string | null, "confidence": "high" | "low",
   "page": int | null, "paragraph": int | null, "clause": string | null, "quote": string,
   "low_reason": {"type": string, "explanation": string, "recommendation": string,
     "statute": {"name": string, "section": string, "url": string} | null,
     "conflict_with": {"contract": string, "detail": string} | null,
     "internal_conflict": {"clause": string, "page": int | null, "paragraph": int | null, "quote": string} | null} | null}]}"""

def build_other_summaries(contracts, exclude_id):
    key_terms = ('Payment Terms', 'Termination', 'Renewal', 'Liability', 'Exclusivity', 'Confidentiality', 'Governing Law')
    summaries = []
    for c in contracts:
        if c.get('id') == exclude_id:
            continue
        facts = ' ; '.join(f"{f['term']}: {f['value'][:120]}" for f in (c.get('facts') or []) if f.get('term') in key_terms)
        summaries.append({
            'name': c.get('contract_name', 'Unknown'),
            'type': c.get('contract_type', 'Agreement'),
            'parties': ', '.join(c.get('parties', [])),
            'dates': f"{c.get('start_date') or '?'} to {c.get('end_date') or '?'}",
            'key_terms': facts[:600]
        })
    return summaries

def ai_extract_matrix(text, filename, others):
    if not OPENROUTER_API_KEY or not text.strip():
        return None
    t = text.strip()
    if len(t) > 40000:
        t = t[:26000] + '\n\n[...middle of document truncated...]\n\n' + t[-10000:]
    if others:
        others_txt = '\n'.join(
            f"- {o['name']} [{o['type']}] parties: {o['parties']} | term: {o['dates']} | key terms: {o['key_terms']}"
            for o in others)
    else:
        others_txt = '(no other contracts)'
    messages = [
        {'role': 'system', 'content': MATRIX_PROMPT.replace('{OTHERS}', others_txt)},
        {'role': 'user', 'content': f'Filename: {filename}\n\nDocument text:\n\n{t}'}
    ]
    try:
        return _parse_json_content(_call_openrouter(messages, max_tokens=4000))
    except Exception:
        print(f'AI matrix extraction failed for {filename}:')
        traceback.print_exc()
        return None

def _norm_matrix_row(r, contract):
    topic_raw = str(r.get('topic') or '').strip()
    key = topic_raw.lower().replace('-', ' ').replace('&', 'and')
    topic = None
    for ct in CANON_TOPICS:
        ck = ct.lower().replace('&', 'and')
        if ck == key or key in ck or ck in key:
            topic = ct
            break
    if topic is None:
        return None

    def val(v):
        if v is None:
            return None
        s = ' '.join(str(v).split())[:400]
        if s.lower() in ('null', 'none', 'n/a'):
            return None
        return s

    conf = str(r.get('confidence') or '').lower().strip()
    if conf not in ('high', 'low'):
        conf = 'high'
    cit = _norm_citation(r.get('page'), r.get('paragraph'), r.get('clause'), r.get('quote'))

    lr = r.get('low_reason') if isinstance(r.get('low_reason'), dict) else None
    low_reason = None
    if conf == 'low':
        ltype = str((lr or {}).get('type') or 'ambiguity').lower().strip()
        if ltype not in ('ambiguity', 'statute', 'conflict', 'internal'):
            ltype = 'ambiguity'
        statute = None
        st = (lr or {}).get('statute')
        if isinstance(st, dict) and st.get('name'):
            url = str(st.get('url') or '').strip()
            try:
                host = urlparse(url).netloc.lower()
            except Exception:
                host = ''
            if not any(host == h or host.endswith('.' + h) for h in ALLOWED_REGULATOR_HOSTS):
                url = 'https://sso.agc.gov.sg/'
            statute = {
                'name': str(st['name'])[:120],
                'section': (str(st.get('section'))[:60] if st.get('section') else None),
                'url': url[:250]
            }
        conflict_with = None
        cw = (lr or {}).get('conflict_with')
        if isinstance(cw, dict) and cw.get('contract'):
            conflict_with = {
                'contract': str(cw['contract'])[:150],
                'detail': ' '.join(str(cw.get('detail') or '').split())[:600]
            }
        internal_conflict = None
        ic = (lr or {}).get('internal_conflict')
        if isinstance(ic, dict) and (ic.get('clause') or ic.get('quote')):
            ic_cit = _norm_citation(ic.get('page'), ic.get('paragraph'), None, ic.get('quote'))
            internal_conflict = {
                'clause': (str(ic.get('clause'))[:100] if ic.get('clause') else None),
                'page': ic_cit['page'],
                'paragraph': ic_cit['paragraph'],
                'quote': ic_cit['quote']
            }
        low_reason = {
            'type': ltype,
            'explanation': ' '.join(str((lr or {}).get('explanation') or 'Contract language requires review.').split())[:900],
            'recommendation': ' '.join(str((lr or {}).get('recommendation') or 'Review and clarify this clause with the counterparty.').split())[:600],
            'statute': statute,
            'conflict_with': conflict_with,
            'internal_conflict': internal_conflict
        }

    return {
        'topic': topic,
        'party1': val(r.get('party1')),
        'party2': val(r.get('party2')),
        'both': val(r.get('both')),
        'confidence': conf,
        'low_reason': low_reason,
        'page': cit['page'],
        'paragraph': cit['paragraph'],
        'clause': cit['clause'],
        'quote': cit['quote']
    }

def _normalize_matrix(raw, contract):
    parties = contract.get('parties') or []
    p1 = raw.get('party1') if isinstance(raw.get('party1'), dict) else {}
    p2 = raw.get('party2') if isinstance(raw.get('party2'), dict) else {}
    party1 = {'name': ' '.join(str(p1.get('name') or (parties[0] if parties else 'Party 1')).split())[:120],
              'role': (' '.join(str(p1.get('role')).split())[:60] if p1.get('role') else None)}
    party2 = {'name': ' '.join(str(p2.get('name') or (parties[1] if len(parties) > 1 else 'Party 2')).split())[:120],
              'role': (' '.join(str(p2.get('role')).split())[:60] if p2.get('role') else None)}

    rows_by_topic = {}
    for r in (raw.get('rows') or []):
        if not isinstance(r, dict):
            continue
        nrm = _norm_matrix_row(r, contract)
        if nrm:
            rows_by_topic[nrm['topic']] = nrm

    rows = []
    for topic in CANON_TOPICS:
        row = rows_by_topic.get(topic)
        if row:
            rows.append(row)
        else:
            rows.append({'topic': topic, 'party1': None, 'party2': None, 'both': None,
                         'confidence': None, 'low_reason': None, 'page': None,
                         'paragraph': None, 'clause': None, 'quote': None, 'not_addressed': True})
    return {'party1': party1, 'party2': party2, 'rows': rows}

def extract_contract_info(text, filename):
    name = os.path.splitext(filename)[0].replace('_', ' ').replace('-', ' ').title()
    
    type_checks = [
        ('Lease', [r'\blease\b', r'\blessor\b', r'\blessee\b', r'\blandlord\b', r'\btenant\b', r'\brent\b']),
        ('NDAs', [r'non-disclosure', r'nondisclosure', r'\bnda\b', r'confidentiality']),
        ('Distribution Agreement', [r'distribution agreement', r'\bdistributor\b', r'channel partner', r'\breseller\b']),
        ('Supplier Contract', [r'supply agreement', r'supplier contract', r'\bsupplier\b', r'\bshipper\b', r'\bcarrier\b', r'\bmanufacturer\b', r'\bbuyer\b']),
        ('Customer Terms', [r'customer terms', r'terms of service', r'subscription', r'\bsubscriber\b', r'\bcustomer\b']),
    ]
    head = text[:600].lower()
    text_lower = text.lower()
    contract_type = 'Agreement'
    for ctype, patterns in type_checks:
        if any(re.search(p, head) for p in patterns):
            contract_type = ctype
            break
    else:
        scores = {ctype: sum(len(re.findall(p, text_lower)) for p in patterns) for ctype, patterns in type_checks}
        best = max(scores, key=scores.get)
        if scores[best] > 0:
            contract_type = best

    parties = []
    paren_match = re.search(
        r'by and between\s+([^(;:]{2,140}?)\s*\(([^)]{1,60})\)\s*(?:and|&)\s+([^(;:]{2,140}?)\s*\(([^)]{1,60})\)',
        text, re.IGNORECASE)
    if paren_match:
        for name in (paren_match.group(1), paren_match.group(3)):
            clean = ' '.join(name.split())
            if clean.lower() not in [p.lower() for p in parties]:
                parties.append(clean)
    if not parties:
        plain_match = re.search(
            r'by and between\s+([^,;:]{2,140}?)\s+and\s+([^,;:]{2,140}?)(?:\s*\(|\.)',
            text, re.IGNORECASE | re.DOTALL)
        if plain_match:
            for name in (plain_match.group(1), plain_match.group(2)):
                clean = ' '.join(name.split())
                if clean.lower() not in [p.lower() for p in parties]:
                    parties.append(clean)

    generic_names = {'the parties', 'parties', 'party a', 'party b', 'both parties', 'the party', 'such party'}
    parties = [p for p in parties if p.lower() not in generic_names and 2 < len(p) < 100]
    
    dates_found = extract_dates(text)
    start_date = None
    end_date = None
    
    for d_str in dates_found:
        parsed = parse_date(d_str)
        if parsed:
            if not start_date:
                start_date = parsed
            elif parsed > start_date and not end_date:
                end_date = parsed
    
    if not start_date and dates_found:
        for d_str in dates_found:
            parsed = parse_date(d_str)
            if parsed:
                start_date = parsed
                break
    
    if not end_date and len(dates_found) >= 2:
        for d_str in dates_found[-1:]:
            parsed = parse_date(d_str)
            if parsed and parsed > start_date:
                end_date = parsed
                break
    
    if not end_date:
        if start_date:
            end_date = start_date.replace(year=start_date.year + 1)
    
    deliverables = []
    deadline_keywords = [
        r'(?:deliverable|milestone|obligation|deadline|due\s*(?:date|by|on))\s*:?\s*(\d{1,2}[\/\-\.]\d{1,2}[\/\-\.]\d{2,4})',
        r'(?:due\s*(?:on|by|date)|deadline|must\s*be\s*(?:completed|submitted|delivered))\s*by\s*(\d{1,2}[\/\-\.]\d{1,2}[\/\-\.]\d{2,4})',
        r'(\d{1,2}\s+\w+\s+\d{4})\s*(?:is\s*(?:the\s+)?(?:deadline|due\s*date|milestone))',
    ]
    for pattern in deadline_keywords:
        matches = re.findall(pattern, text, re.IGNORECASE)
        for m in matches:
            if isinstance(m, tuple):
                for d in m:
                    if d:
                        parsed = parse_date(d)
                        if parsed:
                            deliverables.append(parsed)
            else:
                parsed = parse_date(m)
                if parsed:
                    deliverables.append(parsed)
    
    return {
        'id': str(uuid.uuid4()),
        'filename': filename,
        'contract_name': name,
        'contract_type': contract_type,
        'parties': parties[:3],
        'start_date': start_date.isoformat() if start_date else None,
        'end_date': end_date.isoformat() if end_date else None,
        'deliverables': [d.isoformat() for d in sorted(deliverables)],
        'uploaded_at': dt.now().isoformat(),
        'text_preview': text[:500]
    }

@app.route('/api/upload', methods=['POST'])
def upload():
    if 'file' not in request.files:
        return jsonify({'error': 'No file provided'}), 400

    file = request.files['file']
    orig_name = os.path.basename((file.filename or '').replace('\\', '/'))
    if not orig_name:
        return jsonify({'error': 'No file selected'}), 400

    ext = os.path.splitext(orig_name)[1].lower()
    if ext == '.doc':
        return jsonify({'error': 'Legacy .doc is not supported - please save as .docx or export to PDF'}), 400
    if not allowed_file(orig_name):
        return jsonify({'error': f'Unsupported file type "{ext}" - use PDF, DOCX or TXT'}), 400

    unique_id = str(uuid.uuid4())[:8]
    safe_name = f"{unique_id}_{orig_name}"
    filepath = os.path.join(UPLOAD_FOLDER, safe_name)
    file.save(filepath)

    try:
        contract_info = process_document(filepath, orig_name)
        contract_info['filepath'] = safe_name

        contracts = load_contracts()
        contracts.append(contract_info)
        save_contracts(contracts)

        return jsonify({
            'message': 'File processed successfully',
            'extraction_method': contract_info.get('extraction_method'),
            'contract': contract_info
        }), 200

    except Exception as e:
        traceback.print_exc()
        return jsonify({'error': str(e), 'trace': traceback.format_exc()}), 500

@app.route('/api/upload_multiple', methods=['POST'])
def upload_multiple():
    files = request.files.getlist('files')
    results = []
    errors = []

    for file in files:
        raw_name = (file.filename or '').replace('\\', '/')
        orig_name = os.path.basename(raw_name)
        ext = os.path.splitext(orig_name)[1].lower()
        if not orig_name:
            errors.append({'file': file.filename or '(unnamed)', 'error': 'Missing filename'})
            continue
        if ext == '.doc':
            errors.append({'file': orig_name, 'error': 'Legacy .doc is not supported - please save as .docx or export to PDF'})
            continue
        if not allowed_file(orig_name):
            errors.append({'file': orig_name, 'error': f'Unsupported file type "{ext}" - use PDF, DOCX or TXT'})
            continue

        try:
            unique_id = str(uuid.uuid4())[:8]
            safe_name = f"{unique_id}_{orig_name}"
            filepath = os.path.join(UPLOAD_FOLDER, safe_name)
            file.save(filepath)

            contract_info = process_document(filepath, orig_name)
            contract_info['filepath'] = safe_name

            contracts = load_contracts()
            contracts.append(contract_info)
            save_contracts(contracts)

            results.append(contract_info)
        except Exception as e:
            traceback.print_exc()
            errors.append({'file': file.filename, 'error': str(e)})

    return jsonify({'results': results, 'errors': errors, 'total': len(results)}), 200

@app.route('/api/reprocess', methods=['POST'])
def reprocess_all():
    contracts = load_contracts()
    updated = []
    errors = []
    methods = {'ai': 0, 'ai_vision': 0, 'regex': 0}

    for c in contracts:
        fp = os.path.join(UPLOAD_FOLDER, c.get('filepath', ''))
        if not c.get('filepath') or not os.path.isfile(fp):
            errors.append({'id': c.get('id'), 'file': c.get('filename'), 'error': 'source file missing'})
            updated.append(c)
            continue
        try:
            orig_name = c.get('filename') or os.path.basename(c['filepath'])
            new_info = process_document(fp, orig_name)
            new_info['id'] = c.get('id')
            new_info['uploaded_at'] = c.get('uploaded_at')
            new_info['filepath'] = c.get('filepath')
            updated.append(new_info)
            methods[new_info.get('extraction_method', 'regex')] += 1
        except Exception as e:
            traceback.print_exc()
            errors.append({'id': c.get('id'), 'file': c.get('filename'), 'error': str(e)})
            updated.append(c)

    save_contracts(updated)
    return jsonify({
        'message': f"Reprocessed {len(updated) - len(errors)} of {len(contracts)} contracts",
        'extraction_methods': methods,
        'errors': errors
    }), 200

@app.route('/api/contracts', methods=['GET'])
def get_contracts():
    contracts = load_contracts()
    sorted_contracts = sorted(contracts, key=lambda c: (c.get('end_date') or c.get('start_date') or '9999'), reverse=False)
    return jsonify({'contracts': sorted_contracts, 'total': len(contracts)}), 200

@app.route('/api/contracts/<contract_id>', methods=['GET'])
def get_contract(contract_id):
    contracts = load_contracts()
    contract = next((c for c in contracts if c['id'] == contract_id), None)
    if not contract:
        return jsonify({'error': 'Contract not found'}), 404
    return jsonify({'contract': contract}), 200

@app.route('/api/calendar', methods=['GET'])
def get_calendar():
    contracts = load_contracts()
    events = []
    
    for contract in contracts:
        if contract.get('start_date'):
            events.append({
                'id': contract['id'],
                'title': f"Start: {contract['contract_name']}",
                'start': contract['start_date'],
                'end': contract['start_date'],
                'type': 'start',
                'contract_type': contract.get('contract_type', 'Agreement'),
                'parties': ', '.join(contract.get('parties', [])),
                'color': '#3b82f6'
            })
        if contract.get('end_date'):
            events.append({
                'id': contract['id'],
                'title': f"End: {contract['contract_name']}",
                'start': contract['end_date'],
                'end': contract['end_date'],
                'type': 'end',
                'contract_type': contract.get('contract_type', 'Agreement'),
                'parties': ', '.join(contract.get('parties', [])),
                'color': '#ef4444'
            })
        for dl in contract.get('deliverables', []):
            events.append({
                'id': contract['id'],
                'title': f"Deliverable: {contract['contract_name']}",
                'start': dl,
                'end': dl,
                'type': 'deliverable',
                'contract_type': contract.get('contract_type', 'Agreement'),
                'parties': ', '.join(contract.get('parties', [])),
                'color': '#f59e0b'
            })
    
    events.sort(key=lambda e: e['start'])
    return jsonify({'events': events}), 200

@app.route('/api/delete/<contract_id>', methods=['DELETE'])
def delete_contract(contract_id):
    contracts = load_contracts()
    contracts = [c for c in contracts if c['id'] != contract_id]
    save_contracts(contracts)
    return jsonify({'message': 'Contract deleted'}), 200

@app.route('/api/matrix/<contract_id>', methods=['GET', 'POST'])
def get_matrix(contract_id):
    contracts = load_contracts()
    c = next((x for x in contracts if x.get('id') == contract_id), None)
    if not c:
        return jsonify({'error': 'Contract not found'}), 404
    force = request.method == 'POST' or request.args.get('force') == '1'
    if not force and c.get('matrix'):
        return jsonify({'matrix': c['matrix'], 'cached': True}), 200

    fp = os.path.join(UPLOAD_FOLDER, c.get('filepath', ''))
    if not c.get('filepath') or not os.path.isfile(fp):
        return jsonify({'error': 'Source file missing'}), 400

    ext = os.path.splitext(c.get('filename') or fp)[1].lower()
    if ext == '.pdf':
        text = extract_text_from_pdf(fp)
    elif ext == '.docx':
        text = extract_text_from_docx(fp)
    else:
        with open(fp, 'r', encoding='utf-8', errors='ignore') as f:
            text = f.read()

    others = build_other_summaries(contracts, contract_id)
    raw = ai_extract_matrix(text, c.get('filename') or os.path.basename(fp), others)
    if raw is None:
        return jsonify({'error': 'AI matrix extraction failed. Check API key/connectivity.'}), 502

    matrix = _normalize_matrix(raw, c)
    for row in matrix['rows']:
        lr = row.get('low_reason')
        if lr and lr.get('conflict_with'):
            target_name = lr['conflict_with']['contract'].lower()
            target = next((x for x in contracts if x.get('id') != contract_id and (
                x.get('contract_name', '').lower() == target_name or
                target_name in x.get('contract_name', '').lower() or
                x.get('contract_name', '').lower() in target_name)), None)
            lr['conflict_with']['contract_id'] = target.get('id') if target else None
            lr['conflict_with']['contract_filepath'] = target.get('filepath') if target else None
    matrix['generated_at'] = dt.now().isoformat()
    c['matrix'] = matrix
    save_contracts(contracts)
    return jsonify({'matrix': matrix, 'cached': False}), 200

@app.route('/api/page_image/<filename>/<int:page_no>')
def page_image(filename, page_no):
    safe = os.path.basename(filename)
    path = os.path.join(UPLOAD_FOLDER, safe)
    if not os.path.isfile(path) or not safe.lower().endswith('.pdf'):
        return jsonify({'error': 'File not found'}), 404
    q = (request.args.get('q') or '').strip()
    doc = fitz.open(path)
    try:
        if page_no < 1 or page_no > doc.page_count:
            return jsonify({'error': 'Page out of range'}), 404
        page = doc[page_no - 1]
        rects = []
        if q:
            for attempt in (q, q[:80], q[:50], q[:25]):
                if attempt.strip():
                    try:
                        rects = page.search_for(attempt.strip())
                    except Exception:
                        rects = []
                    if rects:
                        break
        pix = page.get_pixmap(dpi=150)
        img = Image.frombytes('RGB', [pix.width, pix.height], pix.samples)
        if rects:
            overlay = Image.new('RGBA', img.size, (0, 0, 0, 0))
            od = ImageDraw.Draw(overlay)
            scale = 150 / 72
            for r in rects:
                od.rectangle([r.x0 * scale, r.y0 * scale, r.x1 * scale, r.y1 * scale], fill=(255, 214, 0, 105))
            img = Image.alpha_composite(img.convert('RGBA'), overlay).convert('RGB')
        buf = io.BytesIO()
        img.save(buf, format='PNG')
        buf.seek(0)
        return send_file(buf, mimetype='image/png', max_age=3600)
    finally:
        doc.close()

@app.route('/viewer')
def viewer():
    return send_from_directory(os.path.join(os.path.dirname(os.path.dirname(__file__)), 'frontend'), 'viewer.html')

@app.route('/api/stats', methods=['GET'])
def get_stats():
    contracts = load_contracts()
    types = {}
    for c in contracts:
        ct = c.get('contract_type', 'Agreement')
        types[ct] = types.get(ct, 0) + 1
    
    upcoming = []
    for c in contracts:
        if not c.get('end_date'):
            continue
        try:
            ed = dt.fromisoformat(c['end_date'])
        except (ValueError, TypeError):
            ed = parse_date(c['end_date'])
        if ed and ed > dt.now():
            upcoming.append(c)
    
    return jsonify({
        'total': len(contracts),
        'by_type': types,
        'upcoming_deadlines': len(upcoming),
        'with_parties': len([c for c in contracts if len(c.get('parties', [])) > 0])
    }), 200

@app.route('/')
def index():
    return send_from_directory(os.path.join(os.path.dirname(os.path.dirname(__file__)), 'frontend'), 'index.html')

@app.route('/<path:filename>')
def serve_static(filename):
    return send_from_directory(os.path.join(os.path.dirname(os.path.dirname(__file__)), 'frontend'), filename)

@app.route('/api/uploads/<filename>')
def serve_upload(filename):
    return send_from_directory(UPLOAD_FOLDER, filename)

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
